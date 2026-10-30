# Génère un classeur V1.3 fictif pour les tests d'import (python3 generer-classeur-v13.py).
# Toutes les données sont inventées.
import datetime as dt
import openpyxl

D = dt.datetime
wb = openpyxl.Workbook()
wb.remove(wb.active)

def feuille(nom, entetes, lignes):
    ws = wb.create_sheet(nom)
    ws.append(entetes)
    for l in lignes:
        ws.append(l)

feuille('Patients', ['IPP', 'Nom', 'Prenom', 'Naissance', 'Service', "Mode d'hospitalisation", 'Actif', 'Mesure de protection juridique (OUI/NON)', 'Nom du tuteur/curateur'], [
    ['900201', 'DUPONT (fictif)', 'Jean', D(1980, 4, 2), 'UHP', 'SDT', 'OUI', 'NON', None],
])
M = ['ID', 'ID Registre', 'IPP', 'Nom', 'Type', 'Debut', 'Derniere evaluation', 'Psychiatre', 'Actions', 'Duree reelle', 'Reevaluation avant',
     'Info avant', 'Saisine avant', 'Decision avant', 'Fin', 'Tiers', 'Personne', 'Informateur', 'Info effective', 'Saisine effective',
     'Decision effective', 'Cycle', 'Observations', 'Modif par', 'Modif le', 'JLD levee', 'Hors delai', 'Reeval hors delai', 'Ecart reeval',
     'Lieu (CSI/Autre)', 'Nom du lieu', 'Exception CSI (contention)', 'Annee registre']
def mesure(**k):
    return [k.get(c.replace(' ', '_').replace('(', '').replace(')', '').replace('/', '_')) for c in M]
feuille('Mesures', M, [
    mesure(ID='M0001', ID_Registre='A11', IPP='900201', Nom='DUPONT (fictif) Jean', Type='Isolement', Debut=D(2026, 3, 1, 8, 45), Fin=D(2026, 3, 2, 8, 45),
           Psychiatre='Dr Exemple (fictif)', Tiers='OUI', Personne='Mère', Informateur='Dr Exemple (fictif)', Cycle=1, JLD_levee='NON', Lieu_CSI_Autre='CSI', Nom_du_lieu='410'),
    mesure(ID='M0002', IPP='900201', Nom='DUPONT (fictif) Jean', Type='Isolement', Debut=D(2026, 10, 1, 8, 45), Psychiatre='Dr Exemple (fictif)',
           Info_avant=D(2026, 10, 7, 8, 45), Saisine_avant=D(2026, 10, 8, 8, 45), Decision_avant=D(2026, 10, 9, 8, 45), Cycle=2, JLD_levee='NON',
           Lieu_CSI_Autre='CSI', Nom_du_lieu='412'),
    mesure(ID='M0003', ID_Registre='D12', IPP='900202', Nom='MARTIN (fictif) Léa', Type='Contention', Debut=D(2026, 5, 5, 10, 0), Fin=D(2026, 5, 5, 16, 0),
           Psychiatre='Dr Autre (fictif)', Tiers='NON', Informateur='IDE de garde', Cycle=1, JLD_levee='NON', Lieu_CSI_Autre='Autre', Nom_du_lieu='Chambre 12',
           Exception_CSI_contention='OUI'),
    mesure(ID='M0004', Type='Isolement', Debut=D(2026, 6, 1, 9, 0)),
])
feuille('HISTORIQUE_CYCLES_JLD', ['ID Cycle', 'ID Mesure', 'ID Registre', 'IPP', 'Nom', 'Type', 'Cycle', 'Debut', 'Tiers', 'Personne', 'Lien', 'Date tiers',
                                  'Informateur', 'Motif', 'Info avant', 'Info effective', 'Saisine avant', 'Saisine effective', 'Decision avant',
                                  'Decision effective', 'Hors delai', 'Auteur', 'Observations'], [
    ['CJ0001', 'M0002', None, '900201', 'DUPONT (fictif) Jean', 'Isolement', 1, D(2026, 10, 1, 8, 45), 'OUI', 'Sœur', None, D(2026, 10, 1, 10, 0),
     'Dr Exemple (fictif)', None, D(2026, 10, 3, 8, 45), D(2026, 10, 3, 8, 0), D(2026, 10, 4, 8, 45), D(2026, 10, 4, 8, 0), D(2026, 10, 5, 8, 45),
     D(2026, 10, 4, 15, 0), 'NON', 'Secrétaire (fictive)', None],
    ['CJ0002', 'M0002', None, '900201', 'DUPONT (fictif) Jean', 'Isolement', 2, D(2026, 10, 1, 8, 45), 'IMPOSSIBLE A JOINDRE', 'Sœur', None, D(2026, 10, 6, 9, 0),
     'IDE', None, None, None, None, None, None, None, 'NON', None, None],
])
feuille('SUIVI_PROFESSIONNELS_CSI', ['ID', 'Date', 'Periode', 'Service', 'Nom', 'Fonction', 'Portee', 'Lieux (si portee specifique)', 'Saisi par', 'Saisi le'], [
    ['SP1', D(2026, 10, 2), 'Matin', 'UHP', 'Agent Z (fictif)', 'IDE', None, None, None, None],
    ['SP2', D(2026, 10, 2), 'Matin', 'UHP', 'A renseigner', None, None, None, None, None],
])
feuille('CONFIG_SERVICES', ['Code', 'Libelle', 'Actif', 'Defaut'], [['UR', 'Urgences', 'NON', 'NON'], ['UHP', "Unité d'hospitalisation (fictive)", 'OUI', 'OUI']])
feuille('CONFIG_LIEUX_CSI', ['Service', 'Lieu'], [['UHP', 410], ['UHP', 412]])
feuille('CONFIG_PERIODES_SUIVI', ['Service', 'Ordre', 'Periode', 'Heure debut', 'Heure fin'], [['UHP', 1, 'Matin', dt.time(7, 0), dt.time(19, 0)], ['UHP', 2, 'Nuit', dt.time(19, 0), dt.time(7, 0)]])
feuille('CONFIG_COMPTEURS_REGISTRE', ['Service', 'Annee', 'Dernier numero utilise'], [['UHP', 2026, 12]])
feuille('CONFIG_MASQUES_EXPORT', ['Service', 'Feuille', 'Actif', 'Titre isolement', 'Titre contention', 'Sous-titre', 'Afficher sous-titre',
                                  'Masquer N (OUI/NON)', 'Masquer Debut (OUI/NON)', 'Masquer Fin (OUI/NON)', 'Masquer Duree (OUI/NON)', 'Masquer Cumul (OUI/NON)'], [
    ['UHP', 'MASQUE_UHP', 'OUI', "Cumul des mesures d'isolement", 'Cumul des mesures de contention', 'Registre de suivi', 'NON', 'OUI', 'NON', 'NON', 'NON', 'NON'],
])
ws = wb.create_sheet('PARAMETRES')
for l in [('Parametre', 'Isolement', 'Contention'), ('Information JLD - seuil cumul 15j (heures)', 48, 24), ('Seuil jaune avant echeance (heures)', 8, 12)]:
    ws.append(l)
wb.save('classeur-v13-fictif.xlsx')
print('ok')
